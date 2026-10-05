from __future__ import annotations

import json
import re
import shutil
import sys
from collections import Counter
from datetime import datetime
from pathlib import Path

import pdfplumber
from docx import Document
from openpyxl import load_workbook


PROJECT = Path(__file__).resolve().parent
ROOT = PROJECT.parent
PREVIEW_ONLY = "--preview" in sys.argv
DIST = PROJECT / "tmp" / "preview" if PREVIEW_ONLY else PROJECT / "dist"
FILES_OUT = DIST / "files"

HEADERS = [
    "stt", "code", "field", "stages", "group", "name", "indicators",
    "owner", "source_type", "format", "purpose", "frequency",
]

FIELD_MAP = {
    "1": "Ngân sách trung ương của bộ, cơ quan trung ương",
    "2": "Ngân sách địa phương",
    "3": "Quyết toán ngân sách nhà nước",
    "4": "Dự án đầu tư xây dựng",
    "5": "Doanh nghiệp và các tổ chức tài chính, ngân hàng",
    "6": "Chuyên đề",
    "7": "Phục vụ chung toàn Ngành",
}

STAGE_MAP = {
    "A1": "Chuẩn bị kiểm toán",
    "A2": "Thực hiện kiểm toán",
    "A3": "Lập và gửi báo cáo kiểm toán",
    "A4": "Theo dõi thực hiện kết luận, kiến nghị kiểm toán",
    "B1": "Xây dựng kế hoạch kiểm toán trung hạn và hằng năm",
    "B2": "Kiểm soát chất lượng kiểm toán",
    "B3": "Tổng hợp, báo cáo kết quả kiểm toán",
    "B4": "Theo dõi thực hiện kết luận, kiến nghị ở cấp Ngành",
    "B5": "Công khai kết quả kiểm toán",
    "B6": "Thanh tra",
}

SOURCE_TYPE_MAP = {
    "1": "Hệ thống của đơn vị được kiểm toán hoặc đơn vị liên quan",
    "2": "Cơ sở dữ liệu quốc gia, chuyên ngành",
    "3": "Hồ sơ, tài liệu rời do đơn vị quản lý",
    "4": "Dữ liệu từ cơ quan thanh tra, kiểm tra, giám sát",
    "5": "Dữ liệu công khai, dữ liệu mở",
    "6": "Dữ liệu phải mua hoặc thuê quyền truy cập",
}

FORMAT_MAP = {
    "1": "Bản giấy",
    "2": "Bản scan hoặc PDF không bóc tách được",
    "3": "Tệp văn bản hoặc bảng tính rời",
    "4": "Dữ liệu trong hệ thống thông tin, cơ sở dữ liệu",
    "5": "Dữ liệu công bố trên môi trường mạng",
    # The source legend only defines 01-05. Keep an out-of-range entry
    # readable in the UI instead of leaking the raw numeric code into filters.
    "6": "Chưa xác định (mã ngoài danh mục hình thức chuẩn)",
}

FREQUENCY_MAP = {
    "1": "Thường xuyên",
    "2": "Định kỳ",
    "3": "Không thường xuyên",
}

# The 30-unit roster supplied by the user is the baseline for response tracking.
UNIT_ROSTER = [
    ("VP", "Văn phòng Kiểm toán nhà nước", "Tham mưu"),
    ("TCCB", "Vụ Tổ chức cán bộ", "Tham mưu"),
    ("VTH", "Vụ Tổng hợp", "Tham mưu"),
    ("VCS", "Vụ Chính sách kiểm toán nhà nước", "Tham mưu"),
    ("PC", "Vụ Pháp chế", "Tham mưu"),
    ("HTQT", "Vụ Hợp tác quốc tế", "Tham mưu"),
    ("TTRA", "Thanh tra Kiểm toán nhà nước", "Tham mưu"),
    ("VPDU", "Văn phòng Đảng ủy", "Tham mưu"),
    ("CNTT", "Cục Công nghệ thông tin", "Tham mưu"),
    ("CNIa", "Kiểm toán nhà nước chuyên ngành Ia", "Chuyên ngành"),
    ("CNIb", "Kiểm toán nhà nước chuyên ngành Ib", "Chuyên ngành"),
    ("CN2", "Kiểm toán nhà nước chuyên ngành II", "Chuyên ngành"),
    ("CN3", "Kiểm toán nhà nước chuyên ngành III", "Chuyên ngành"),
    ("CN4", "Kiểm toán nhà nước chuyên ngành IV", "Chuyên ngành"),
    ("CN5", "Kiểm toán nhà nước chuyên ngành V", "Chuyên ngành"),
    ("CN6", "Kiểm toán nhà nước chuyên ngành VI", "Chuyên ngành"),
    *[(f"KV{i}", f"Kiểm toán nhà nước khu vực {roman}", "Khu vực")
      for i, roman in enumerate(["I", "II", "III", "IV", "V", "VI", "VII", "VIII", "IX", "X", "XI", "XII"], 1)],
    ("TRUONG", "Trường Đào tạo và Bồi dưỡng nghiệp vụ kiểm toán", "Sự nghiệp"),
    ("BAO", "Báo Kiểm toán", "Sự nghiệp"),
]


def clean(value):
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def expand_code_list(value, mapping, alpha=False):
    # Excel may coerce an entry such as "01, 02, 04" to 01/02/2004 while
    # retaining a display format that visually resembles the intended list.
    if isinstance(value, datetime) and not alpha:
        date_codes = [str(value.day), str(value.month), str(value.year % 100)]
        if all(code in mapping for code in date_codes):
            return "; ".join(dict.fromkeys(mapping[code] for code in date_codes))

    text = clean(value)
    if not text:
        return ""

    # PDF extraction can split a zero-padded code at the comma glyph
    # (for example "0,3" instead of "03"). Code 0 is not valid in any of
    # these classifications, so joining it back is unambiguous.
    if not alpha:
        text = re.sub(r"\b0\s*,\s*([1-9])\b", r"0\1", text)

    # Some units append a free-text subject to the field code. Handle the
    # complete cell before splitting because the subject itself may contain
    # commas (for example: "6 - Quản lý, sử dụng đất đai").
    if mapping is FIELD_MAP:
        field_with_note = re.fullmatch(r"([0-9\s,;/.+]+)\s*[-–:]\s*(.+)", text)
        if field_with_note:
            codes = re.findall(r"\d+", field_with_note.group(1))
            keys = [code.lstrip('0') or '0' for code in codes]
            if all(key in mapping for key in keys):
                labels = list(dict.fromkeys(mapping[key] for key in keys))
                labels[-1] += f": {field_with_note.group(2)}"
                return "; ".join(labels)

    # Excel sometimes turns lists such as "2.3" or "1.2" into decimal-like
    # values. In these classification columns the dot is a separator, not a
    # decimal point. Only split cells that consist entirely of valid codes so
    # ordinary prose containing punctuation is preserved verbatim.
    token = r"[A-Za-z]\d+" if alpha else r"\d+"
    separator = r"(?:\s*[,;/.+]\s*|\s+và\s+|\s+)"
    if not re.fullmatch(fr"{token}(?:{separator}{token})*", text, flags=re.IGNORECASE):
        return text

    parts = [part.strip() for part in re.split(separator, text, flags=re.IGNORECASE) if part.strip()]
    expanded = []
    for part in parts:
        key = part.upper() if alpha else part.lstrip("0") or "0"
        if key in mapping:
            expanded.append(mapping[key])
            continue
        return text
    return "; ".join(dict.fromkeys(expanded))


def normalize_frequency(value):
    if isinstance(value, datetime):
        return expand_code_list(value, FREQUENCY_MAP)
    text = clean(value)
    if not text:
        return ""
    if re.fullmatch(r"[0-9\s,;/.+]+", text):
        return expand_code_list(text, FREQUENCY_MAP)
    lowered = text.casefold()
    levels = []
    if "không thường xuyên" in lowered:
        levels.append("Không thường xuyên")
        lowered = lowered.replace("không thường xuyên", "")
    if "thường xuyên" in lowered:
        levels.append("Thường xuyên")
    if "định kỳ" in lowered or "hằng năm" in lowered or "hàng năm" in lowered:
        levels.append("Định kỳ")
    return "; ".join(dict.fromkeys(levels)) or text


def to_record(values, unit):
    raw_values = list(values)[:12]
    raw_values += [""] * (12 - len(raw_values))
    record = dict(zip(HEADERS, [clean(value) for value in raw_values]))
    record["code"] = re.sub(r"\s*-\s*", "-", record["code"])
    record["field"] = expand_code_list(raw_values[2], FIELD_MAP)
    record["stages"] = expand_code_list(raw_values[3], STAGE_MAP, alpha=True)
    record["source_type"] = expand_code_list(raw_values[8], SOURCE_TYPE_MAP)
    record["format"] = expand_code_list(raw_values[9], FORMAT_MAP)
    record["frequency"] = normalize_frequency(raw_values[11])
    record["unit"] = unit
    return record


def is_column_index_row(values, count):
    normalized = [clean(value) for value in list(values)[:count]]
    return normalized == [str(index) for index in range(1, count + 1)]


def content_key(value):
    return re.sub(r"[^\w]+", " ", clean(value).casefold(), flags=re.UNICODE).strip()


def resolve_formula(wb, value):
    if not isinstance(value, str) or not value.startswith("="):
        return value
    match = re.fullmatch(r"='([^']+)'!([A-Z]+\d+)", value)
    if not match:
        match = re.fullmatch(r"=([^!]+)!([A-Z]+\d+)", value)
    if match:
        return wb[match.group(1)][match.group(2)].value
    return value


def read_excel(path, unit, detail_sheet, detail_start, priority_sheet, priority_start):
    wb = load_workbook(path, data_only=False, read_only=False)
    ws = wb[detail_sheet]
    details = []
    for row in ws.iter_rows(min_row=detail_start, max_col=12, values_only=True):
        marker = row[0]
        is_data_marker = isinstance(marker, (int, float)) or (isinstance(marker, str) and marker.startswith("="))
        if is_data_marker and not is_column_index_row(row, 12) and clean(row[1]):
            details.append(to_record(row, unit))

    ws = wb[priority_sheet]
    priorities = []
    rank = 0
    for row in ws.iter_rows(min_row=priority_start, max_col=4, values_only=True):
        marker = row[0]
        if not (isinstance(marker, (int, float)) or (isinstance(marker, str) and marker.startswith("="))):
            continue
        if is_column_index_row(row, 4):
            continue
        code = clean(resolve_formula(wb, row[1]))
        name = clean(resolve_formula(wb, row[2]))
        if not code or not name:
            continue
        rank += 1
        priorities.append({"rank": rank, "code": code, "name": name, "reason": clean(row[3]), "unit": unit})
    return details, priorities


def read_kv4(path):
    detail_pages = {2, 3, 4, 7, 8, 11, 12, 15}
    priority_pages = {5, 9, 13, 16}
    details, priorities = [], []
    with pdfplumber.open(path) as pdf:
        for page_no, page in enumerate(pdf.pages, 1):
            tables = page.extract_tables() or []
            for table in tables:
                if page_no in detail_pages:
                    for row in table:
                        if row and re.fullmatch(r"\d{1,2}", clean(row[0])) and clean(row[1]) not in {"2", "Mã nguồn"}:
                            details.append(to_record(row, "KV4"))
                elif page_no in priority_pages:
                    for row in table:
                        if row and re.fullmatch(r"\d+", clean(row[0])):
                            priorities.append({
                                "rank": len(priorities) + 1,
                                "section_rank": int(clean(row[0])),
                                "code": clean(row[1]),
                                "name": clean(row[2]),
                                "reason": clean(row[3]),
                                "unit": "KV4",
                            })
    return details, priorities


def read_pdf_forms(path, unit, detail_pages, priority_pages, merge_continuations=False):
    details, priorities = [], []
    raw_details = []
    with pdfplumber.open(path) as pdf:
        for page_no in detail_pages:
            for table in pdf.pages[page_no - 1].extract_tables() or []:
                for row in table:
                    if (
                        row
                        and len(row) >= 12
                        and re.fullmatch(r"\d+", clean(row[0]))
                        and clean(row[1]) not in {"", "2", "Mã nguồn"}
                    ):
                        raw_details.append(list(row[:12]))
                    elif merge_continuations and raw_details and row and len(row) >= 12 and not clean(row[0]):
                        for index, value in enumerate(row[:12]):
                            continuation = clean(value)
                            if continuation:
                                raw_details[-1][index] = clean(f"{raw_details[-1][index] or ''} {continuation}")
        for page_no in priority_pages:
            for table in pdf.pages[page_no - 1].extract_tables() or []:
                for row in table:
                    if (
                        row
                        and len(row) >= 4
                        and re.fullmatch(r"\d+", clean(row[0]))
                        and clean(row[1]) not in {"", "2", "Mã nguồn"}
                        and clean(row[2]) not in {"", "3", "Tên dữ liệu"}
                    ):
                        priorities.append({
                            "rank": int(clean(row[0])),
                            "code": clean(row[1]),
                            "name": clean(row[2]),
                            "reason": clean(row[3]),
                            "unit": unit,
                        })
    details = [to_record(row, unit) for row in raw_details]
    return details, priorities


def read_docx_forms(path, unit, detail_table=0, priority_table=1):
    doc = Document(path)
    details, priorities = [], []
    if len(doc.tables) > detail_table:
        for row in doc.tables[detail_table].rows[2:]:
            values = [cell.text for cell in row.cells]
            if values and re.fullmatch(r"\d+", clean(values[0])) and not is_column_index_row(values, 12) and clean(values[1]):
                details.append(to_record(values, unit))
    if len(doc.tables) > priority_table:
        for row in doc.tables[priority_table].rows[1:]:
            values = [clean(cell.text) for cell in row.cells]
            if values and re.fullmatch(r"\d+", values[0]) and not is_column_index_row(values, 4) and values[1] and values[2]:
                priorities.append({
                    "rank": int(values[0]),
                    "code": values[1],
                    "name": values[2],
                    "reason": values[3],
                    "unit": unit,
                })
    return details, priorities


def read_kv6_scan():
    """Transcription of the image-only 17-page signed PDF received from KV VI."""
    rows = [
        ("01", "ĐT-01", "4", "A1, A2", "Chủ trương đầu tư", "Quyết định chủ trương đầu tư và hồ sơ đề xuất chủ trương đầu tư", "Mục tiêu, quy mô, địa điểm; sơ bộ tổng mức đầu tư; cơ cấu nguồn vốn; thời gian, tiến độ thực hiện; cấp quyết định", "Cấp quyết định chủ trương đầu tư; chủ đầu tư", "3", "1, 2", "Xác định mục tiêu, quy mô được phép đầu tư ban đầu và kiểm tra thẩm quyền, tổng mức đầu tư", "1"),
        ("02", "ĐT-01", "4", "A1, A2", "Chủ trương và quyết định đầu tư", "Báo cáo nghiên cứu khả thi; quyết định phê duyệt dự án và các quyết định điều chỉnh", "Tổng mức đầu tư và cơ cấu chi phí; quy mô, giải pháp kỹ thuật chủ yếu; tiến độ; nội dung, lý do và giá trị từng lần điều chỉnh", "Người quyết định đầu tư; chủ đầu tư; cơ quan tham mưu", "3", "1, 2", "Xác định tổng mức đầu tư được duyệt, tính hợp lý và đúng thẩm quyền của các lần điều chỉnh", "1"),
        ("03", "ĐT-02", "4", "A1, A2", "Kế hoạch vốn đầu tư công", "Kế hoạch đầu tư công trung hạn, hằng năm và các quyết định giao, điều chỉnh kế hoạch vốn", "Vốn bố trí theo từng năm; tổng nguồn vốn; số điều chỉnh; kéo dài thời gian thực hiện; giải ngân và giải ngân lũy kế", "Sở Tài chính; chủ đầu tư", "2, 3", "2, 3, 4", "Đối chiếu số vốn thanh toán với kế hoạch được giao và đánh giá tỷ lệ giải ngân", "1"),
        ("04", "BS-01", "4", "A1", "Giám sát, đánh giá đầu tư", "Báo cáo giám sát, đánh giá tổng thể đầu tư của chủ đầu tư và cơ quan quản lý nhà nước về đầu tư công", "Tình hình thực hiện, tiến độ; khó khăn, vướng mắc; danh sách dự án chậm tiến độ, phải điều chỉnh hoặc có vi phạm", "Cơ quan quản lý nhà nước về đầu tư công; chủ đầu tư; Hệ thống thông tin giám sát, đánh giá đầu tư", "2, 3", "3, 4", "Khảo sát, đánh giá rủi ro để lựa chọn dự án và nội dung kiểm toán trọng yếu", "2"),
        ("05", "ĐT-03", "4", "A1, A2", "Thiết kế, dự toán", "Hồ sơ khảo sát xây dựng và hồ sơ thiết kế, dự toán xây dựng công trình", "Quy mô, kết cấu, vật liệu chủ yếu, khối lượng thiết kế và các thay đổi thiết kế giữa các bước", "Chủ đầu tư", "3", "1, 2, 3", "Kiểm tra sự phù hợp giữa khối lượng dự toán, thanh toán với bản vẽ thiết kế", "1"),
        ("06", "ĐT-03", "4", "A2", "Thiết kế, dự toán", "Báo cáo kết quả thẩm định và quyết định phê duyệt thiết kế, dự toán", "Giá trị trước và sau thẩm định; nội dung cắt giảm, điều chỉnh; cơ quan, người thẩm định và phê duyệt", "Chủ đầu tư; cơ quan chuyên môn về xây dựng", "3", "1, 2", "Đánh giá chất lượng và trách nhiệm của công tác thẩm định", "1"),
        ("07", "ĐT-04", "4", "A1, A2", "Lựa chọn nhà thầu", "Kế hoạch lựa chọn nhà thầu; hồ sơ mời thầu, báo cáo đánh giá hồ sơ dự thầu và quyết định phê duyệt kết quả", "Hình thức, phương thức lựa chọn; giá gói thầu; tiêu chuẩn đánh giá; danh sách nhà thầu; giá dự thầu và tỷ lệ tiết kiệm", "Chủ đầu tư; Hệ thống mạng đấu thầu quốc gia", "1, 5", "4, 5", "Kiểm tra tuân thủ quy trình lựa chọn nhà thầu và phát hiện dấu hiệu hạn chế cạnh tranh, chia nhỏ gói thầu", "1"),
        ("08", "ĐT-05", "4", "A2", "Hợp đồng", "Hợp đồng xây lắp, tư vấn, cung cấp thiết bị và các phụ lục hợp đồng", "Loại hợp đồng, giá hợp đồng; điều kiện tạm ứng, thanh toán, điều chỉnh giá; tiến độ và bảo hành", "Chủ đầu tư", "3", "1, 2", "Kiểm tra áp dụng loại hợp đồng, điều khoản điều chỉnh giá và khối lượng phát sinh", "1"),
        ("09", "ĐT-06", "4", "A2", "Nghiệm thu, thanh toán", "Hồ sơ nghiệm thu, thanh toán khối lượng hoàn thành theo từng đợt; hồ sơ quản lý chất lượng công trình", "Biên bản nghiệm thu; bảng tính giá trị khối lượng hoàn thành; khối lượng phát sinh; chứng từ tạm ứng và thanh toán", "Chủ đầu tư", "3", "1, 2", "Kiểm tra khối lượng, đơn giá nghiệm thu, thanh toán so với thiết kế và hợp đồng", "1"),
        ("10", "NS-03", "4", "A2, A4", "Chấp hành ngân sách nhà nước", "Dữ liệu kiểm soát chi, thanh toán vốn đầu tư của dự án qua Kho bạc Nhà nước", "Số tạm ứng, thu hồi tạm ứng, số dư tạm ứng và các khoản nộp trả ngân sách", "Kho bạc Nhà nước", "2", "4", "Đối chiếu giải ngân với hồ sơ chủ đầu tư và phát hiện tạm ứng quá hạn, giảm trừ thanh toán theo kiến nghị", "1"),
        ("11", "ĐT-07", "4", "A2", "Quyết toán dự án", "Báo cáo quyết toán dự án hoàn thành, báo cáo thẩm tra và quyết định phê duyệt quyết toán", "Giá trị đề nghị và được duyệt theo cơ cấu chi phí; tài sản hình thành; công nợ; chi phí không tính vào giá trị tài sản", "Chủ đầu tư; cơ quan tài chính", "3", "1, 2, 3", "Kiểm tra giá trị quyết toán với thanh toán, dự toán và xác định sai chế độ, giá trị tài sản bàn giao", "1"),
        ("12", "ĐT-08", "4", "A2", "Định mức, đơn giá", "Giá vật liệu xây dựng, đơn giá nhân công, giá ca máy và thiết bị thi công công bố tại địa phương", "Giá theo chủng loại, quy cách, địa điểm cung cấp và thời điểm công bố", "Sở Xây dựng", "5", "2, 3, 5", "Đối chiếu giá vật liệu, nhân công, máy thi công trong dự toán, thanh toán với giá công bố", "1"),
        ("13", "ĐT-08", "4", "A2", "Định mức, đơn giá", "Chỉ số giá xây dựng", "Chỉ số giá theo loại công trình, theo yếu tố chi phí, theo địa bàn và kỳ công bố", "Bộ Xây dựng; Sở Xây dựng", "5", "2, 3, 5", "Kiểm tra tính điều chỉnh giá hợp đồng và chi phí dự phòng do trượt giá", "1"),
        ("01", "NS-02", "2", "A1, A2", "Dự toán NSNN", "Nghị quyết HĐND và Quyết định UBND về giao, điều chỉnh dự toán NSĐP", "Tổng dự toán thu, chi NSĐP; phân bổ dự toán cho từng đơn vị, cấp ngân sách; bổ sung có mục tiêu", "HĐND, UBND tỉnh/xã; Sở Tài chính", "1, 2", "1, 2, 3", "Đánh giá tuân thủ trong lập, giao, điều chỉnh dự toán và đối chiếu với khả năng cân đối", "1"),
        ("02", "NS-01", "2", "A1, A2", "Dự toán NSNN", "Văn bản giao dự toán và điều chỉnh, bổ sung dự toán của đơn vị dự toán cấp I", "Chi tiết dự toán giao cho đơn vị trực thuộc theo mục lục NSNN và tính chất chi", "Cơ quan tài chính; đơn vị dự toán cấp I", "3", "1, 2, 3", "Kiểm tra tính chính xác, đúng thẩm quyền và phân bổ dự toán đúng định mức, tiêu chuẩn", "1"),
        ("03", "NS-03", "2", "A1, A2, A4", "Chấp hành NSNN", "Dữ liệu chi NSNN qua Kho bạc Nhà nước", "Chi tiết số kiểm soát chi, thanh toán, tạm ứng, ứng trước dự toán, chuyển nguồn, kết dư ngân sách", "Kho bạc Nhà nước khu vực", "1, 2", "4", "Phân tích xu hướng chi, đối chiếu thực tế với dự toán và phát hiện tạm ứng kéo dài, chuyển nguồn sai quy định", "1"),
        ("04", "NS-04", "2", "A1, A2", "Chấp hành NSNN", "Dữ liệu thu ngân sách theo sắc thuế, người nộp thuế, nợ đọng, miễn, giảm, gia hạn", "Số thuế theo sắc thuế, người nộp thuế; số nợ, nợ phân theo nhóm; tiền thuế được gia hạn, miễn, giảm, hoàn", "Cơ quan thuế", "2", "4", "Đánh giá quản lý thu, hạch toán thu NSNN, chính sách ưu đãi thuế và phát hiện thất thoát thu", "1"),
        ("05", "NS-08", "2", "A2", "Kế toán đơn vị", "Sổ kế toán, BCTC, báo cáo quyết toán kinh phí của đơn vị dự toán hoặc đơn vị sự nghiệp", "Bảng cân đối tài khoản, BCTC, số liệu thu chi, quỹ phát triển hoạt động sự nghiệp và nguồn CCTL", "Đơn vị được kiểm toán", "3", "3, 4", "Kiểm tra hợp lệ của chứng từ chi và xác định số dư kinh phí chuyển nguồn, trích lập các quỹ", "1"),
        ("06", "NS-09", "2", "A2", "Hóa đơn, chứng từ", "Dữ liệu hóa đơn điện tử", "Thông tin hóa đơn, người bán, người mua, danh mục hàng hóa, số lượng, giá trị, thuế GTGT và thời điểm", "Cơ quan thuế", "2", "4", "Đối chiếu hóa đơn điện tử với chứng từ thanh toán và phát hiện hóa đơn rủi ro", "2"),
        ("07", "NS-10", "2", "A2", "Tài sản công", "Dữ liệu đăng ký, quản lý, sử dụng tài sản công", "Danh mục nhà, đất, xe ô tô, trang thiết bị; nguyên giá, giá trị còn lại; hiện trạng sử dụng", "Sở Tài chính; đơn vị quản lý tài sản công", "2", "4", "Đánh giá chấp hành tiêu chuẩn, định mức và phát hiện sử dụng tài sản công sai mục đích", "1"),
        ("08", "NS-11", "2", "A1, A2", "Nợ công", "Dữ liệu vay, trả nợ chính quyền địa phương", "Dư nợ đầu kỳ, số vay trong kỳ, trả nợ gốc/lãi, dư nợ cuối kỳ; vay lại vốn vay nước ngoài của Chính phủ", "Sở Tài chính", "1, 2", "3, 4", "Đánh giá an toàn nợ chính quyền địa phương và tuân thủ hạn mức dư nợ vay", "1"),
        ("09", "NS-12", "2", "A2", "Biên chế, quỹ lương", "Dữ liệu biên chế và quỹ tiền lương", "Số biên chế được giao, biên chế có mặt; quỹ lương được duyệt, quyết toán quỹ lương và đóng BHXH", "Sở Nội vụ; đơn vị dự toán", "1, 3", "3, 4", "Kiểm tra chấp hành chỉ tiêu biên chế, quyết toán quỹ tiền lương và các khoản phụ cấp", "1"),
        ("10", "CĐ-03", "2", "A1, A2", "Đất đai", "Dữ liệu giao đất, cho thuê đất, chuyển mục đích sử dụng đất, xác định giá đất", "Dự án hoặc thửa đất; diện tích; mục đích; quyết định giao, thuê, chuyển mục đích; giá đất và nghĩa vụ tài chính", "Sở Nông nghiệp và Môi trường; Sở Tài chính", "2", "3, 4", "Đánh giá quản lý thu từ đất, xác định nghĩa vụ tài chính và phát hiện nợ đọng, thất thoát tiền sử dụng đất", "1"),
        ("11", "CĐ-02", "2", "A1, A2", "Chương trình mục tiêu quốc gia", "Dữ liệu phân bổ, giải ngân vốn các Chương trình mục tiêu quốc gia", "Kế hoạch vốn, danh mục dự án, số vốn giải ngân, danh sách đối tượng thụ hưởng", "Sở Tài chính", "2", "3, 4", "Kiểm tra tuân thủ, mục tiêu và hiệu quả sử dụng kinh phí chương trình", "2"),
        ("12", "NS-06", "2", "A1, B1, B3", "Quyết toán ngân sách nhà nước", "Báo cáo quyết toán ngân sách địa phương", "Tổng quyết toán thu, chi NSĐP; kết dư ngân sách; số chuyển nguồn và chi tiết từng nguồn kinh phí", "Cơ quan tài chính", "2, 3", "2, 4", "Đối chiếu tổng hợp quyết toán NSĐP; phục vụ lập kế hoạch và tổng hợp kết quả kiểm toán", "2"),
        ("01", "NS-01", "1, 2", "A1, A2", "Dự toán thu NSNN", "Văn bản giao dự toán thu và các lần điều chỉnh, bổ sung", "Tổng dự toán thu; chi tiết theo khoản thu, sắc thuế, cơ quan thu, địa bàn; số điều chỉnh, bổ sung", "Cơ quan tài chính; đơn vị dự toán cấp I", "3", "3", "Đối chiếu dự toán với thực hiện và xác định khoản thu tăng, giảm lớn, điều chỉnh bất thường", "1"),
        ("02", "NS-06", "3", "A1, A2, A3", "Quyết toán thu NSNN", "Báo cáo quyết toán thu NSNN của bộ, cơ quan trung ương và địa phương", "Tổng thu; cơ cấu theo khoản thu; số thu theo cấp ngân sách; hoàn trả, điều chỉnh; số thực hiện so với dự toán", "Cơ quan tài chính", "2, 3", "3, 4", "Đối chiếu số thu quyết toán với dữ liệu cơ quan thu và Kho bạc; phân tích cơ cấu, biến động và rủi ro thu", "1"),
        ("03", "NS-04", "1, 2", "A1, A2", "Quản lý thuế nội địa", "Dữ liệu đăng ký, kê khai và xác định nghĩa vụ thuế của người nộp thuế", "MST, trạng thái NNT; kỳ thuế; sắc thuế; doanh thu hoặc thu nhập chịu thuế; số phải nộp, đã nộp; điều chỉnh, khai bổ sung", "Cơ quan thuế", "2", "4", "Xác định đầy đủ quan hệ NNT; đối chiếu kê khai và nộp thuế; sàng lọc rủi ro thất thu", "1"),
        ("04", "NS-04", "1, 2", "A1, A2", "Quản lý nợ thuế", "Dữ liệu nợ thuế, tiền chậm nộp và cưỡng chế nợ thuế", "Nợ đầu kỳ, phát sinh, đã thu; nợ cuối kỳ, tuổi nợ, tiền chậm nộp; quyết định, biện pháp cưỡng chế và kết quả thu hồi", "Cơ quan thuế", "2", "4", "Đánh giá quản lý nợ, phát hiện nợ kéo dài, chậm áp dụng biện pháp cưỡng chế hoặc số nợ rủi ro không thu hồi", "1"),
        ("05", "NS-04", "1, 2", "A1, A2", "Ưu đãi và xử lý nghĩa vụ thuế", "Dữ liệu miễn, giảm, gia hạn, khoanh hoặc xóa nợ, hoàn thuế và kết quả thanh tra, kiểm tra thuế", "NNT, căn cứ, loại thuế, số tiền, thời gian; quyết định; số hoàn, miễn, giảm; số truy thu, xử phạt, tiền chậm nộp", "Cơ quan thuế", "2", "4", "Kiểm tra đúng đối tượng, điều kiện, thẩm quyền và xác định tác động tài chính, rủi ro cần kiểm tra sau", "1"),
        ("06", "NS-09", "1, 2", "A1, A2", "Hóa đơn, chứng từ", "Dữ liệu hóa đơn điện tử ở cấp hóa đơn và thông tin hàng hóa, dịch vụ", "MST người bán, mua; ngày hóa đơn; doanh thu; hàng hóa, dịch vụ; số lượng; đơn giá; thuế suất; tiền thuế; trạng thái điều chỉnh, thay thế, hủy", "Cơ quan thuế", "2", "4", "Đối chiếu doanh thu hóa đơn với kê khai thuế và phát hiện chênh lệch doanh thu, hóa đơn bất thường, chuỗi giao dịch rủi ro", "1"),
        ("07", "NS-05", "1", "A1, A2", "Thu NSNN từ xuất khẩu, nhập khẩu", "Dữ liệu tờ khai hải quan và nghĩa vụ thuế phát sinh từ hoạt động xuất nhập khẩu", "Doanh nghiệp, số tờ khai, mã HS, hàng hóa, trị giá, xuất xứ, thuế suất, thuế phải nộp, miễn, giảm, hoàn, nợ và ấn định thuế", "Cơ quan hải quan", "2", "4", "Kiểm tra xác định và thu đúng, đủ thuế xuất nhập khẩu; phân tích rủi ro trị giá, mã HS, xuất xứ và nợ thuế", "1"),
        ("08", "CĐ-03", "6 - Quản lý, sử dụng đất đai", "A1, A2", "Nghĩa vụ tài chính về đất", "Dữ liệu giao đất, cho thuê đất, chuyển mục đích và xác định nghĩa vụ tài chính về đất", "Dự án, thửa đất; diện tích; mục đích; quyết định; thời điểm xác định nghĩa vụ; tiền sử dụng đất, thuê đất phải nộp, đã nộp, còn nợ, miễn giảm", "Cơ quan quản lý đất đai; cơ quan thuế", "2", "4", "Đối chiếu quá trình phát sinh và thực hiện nghĩa vụ tài chính; phát hiện chậm xác định, thiếu hoặc chưa thu đủ", "1"),
        ("09", "BS-01", "1, 2, 3", "A1, A2, A3", "Thu NSNN hạch toán tại KBNN", "Dữ liệu thu, hoàn trả, điều chỉnh và hạch toán NSNN tại Kho bạc Nhà nước", "MST hoặc NNT; ngày nộp, ngày hạch toán; số chứng từ, số tiền; chương, tiểu mục; cơ quan thu; địa bàn; cấp ngân sách hưởng; tỷ lệ điều tiết; hoàn trả, điều chỉnh", "Kho bạc Nhà nước", "2", "4", "Đối chiếu độc lập số phải nộp với số đã nộp, đã hạch toán; xác minh chênh lệch và khoản thu không phụ thuộc báo cáo cơ quan thu", "1"),
        ("01", "DN-01", "5", "A1, A2", "Đăng ký doanh nghiệp", "Dữ liệu đăng ký doanh nghiệp, vốn điều lệ, người đại diện theo pháp luật", "Vốn đăng ký, vốn thực góp, người đại diện theo pháp luật", "Sở Tài chính", "5", "5", "Đối chiếu tư cách pháp lý, thông tin đăng ký và tình hình góp vốn để xác định phạm vi, đối tượng, rủi ro kiểm toán", "1"),
        ("02", "DN-02", "5", "A1, A2", "Báo cáo tài chính doanh nghiệp", "Báo cáo tài chính riêng, báo cáo tài chính hợp nhất và báo cáo kiểm toán độc lập", "Các chỉ tiêu trên báo cáo tài chính", "Đơn vị được kiểm toán; cơ quan quản lý đơn vị; cơ quan thuế", "5", "5", "Phân tích tình hình tài chính, kết quả hoạt động, đối chiếu số liệu và nhận diện khoản mục rủi ro trọng yếu", "1"),
        ("03", "DN-03", "5", "A1, A2", "Vốn nhà nước tại doanh nghiệp", "Dữ liệu vốn nhà nước đầu tư tại doanh nghiệp; cổ tức, lợi nhuận nộp ngân sách", "Giá trị vốn nhà nước đầu tư; cổ tức, lợi nhuận phải nộp và đã nộp ngân sách", "Đơn vị được kiểm toán; cơ quan quản lý; cơ quan tài chính", "5", "5", "Kiểm tra quản lý, bảo toàn vốn nhà nước và xác định đầy đủ cổ tức, lợi nhuận phải nộp ngân sách", "1"),
        ("04", "DN-04", "5", "A1, A2", "Công bố thông tin", "Thông tin công bố của doanh nghiệp niêm yết, công ty đại chúng", "Nội dung, thời điểm công bố; thông tin tài chính và quản trị có liên quan", "Cơ quan quản lý đơn vị được kiểm toán", "5", "5", "Đối chiếu thông tin công bố với hồ sơ, báo cáo doanh nghiệp và phát hiện chênh lệch hoặc thiếu thông tin", "1"),
        ("05", "DN-05", "5", "A1, A2", "Hoạt động tín dụng", "Dữ liệu dư nợ, phân loại nợ, trích lập và sử dụng dự phòng rủi ro", "Dư nợ, nhóm nợ, số dự phòng phải trích, đã trích và sử dụng dự phòng", "Cơ quan quản lý đơn vị được kiểm toán", "5", "5", "Đánh giá chất lượng tín dụng, phân loại nợ và trích lập, sử dụng dự phòng rủi ro", "1"),
        ("06", "DN-06", "5", "A1, A2", "Nghĩa vụ thuế của doanh nghiệp", "Dữ liệu kê khai, nộp, hoàn thuế của doanh nghiệp", "Số kê khai, số phải nộp, số đã nộp, số được hoàn; nghĩa vụ còn phải thực hiện", "Cơ quan quản lý đơn vị được kiểm toán; cơ quan thuế", "5", "5", "Đối chiếu nghĩa vụ thuế với số đã nộp, số hoàn; xác định khoản kê khai, nộp thuế chưa đầy đủ hoặc có rủi ro", "1"),
    ]
    details = [to_record(row, "KV6") for row in rows]
    priority_specs = [
        ("ĐT-06", "Hồ sơ nghiệm thu, thanh toán khối lượng hoàn thành", "Cần dữ liệu số để kiểm tra toàn bộ khối lượng, đơn giá và hạn chế chọn mẫu thủ công."),
        ("ĐT-03", "Dự toán xây dựng công trình, dự toán gói thầu", "Cần tệp gốc để chạy hàm chuyển đổi giá trị và đối chiếu mã định mức tự động."),
        ("ĐT-04", "Kết quả lựa chọn nhà thầu", "Kết nối dữ liệu đấu thầu giúp so sánh giá trúng thầu và phát hiện dấu hiệu thông thầu."),
        ("ĐT-08", "Giá vật liệu xây dựng do địa phương công bố", "Số hóa thành thư viện dùng chung để tự động hóa khảo sát và phát hiện giá bất thường."),
        ("NS-03", "Dữ liệu chi NSNN qua Kho bạc Nhà nước", "Nguồn dữ liệu gốc chi tiết từng giao dịch, cần khai thác trực tiếp để phân tích diện rộng."),
        ("NS-04", "Dữ liệu thu NSNN theo sắc thuế, người nộp thuế, nợ thuế", "Khai thác dữ liệu số giúp kiểm tra chính xác người nộp thuế và hoàn thuế."),
        ("CĐ-03", "Dữ liệu giao đất, cho thuê đất, xác định giá đất và tiền sử dụng đất", "Nguồn thu từ đất có tỷ trọng lớn và rủi ro thất thoát cao, cần liên thông để đối chiếu."),
        ("NS-10", "Dữ liệu đăng ký, quản lý, sử dụng tài sản công", "Kết nối dữ liệu tập trung giúp đánh giá hiện trạng sử dụng tài sản công toàn địa bàn."),
        ("NS-02", "Dự toán và phân bổ dự toán NSĐP", "Cần tệp bảng tính gốc để giảm nhập liệu thủ công và đối chiếu phân bổ dự toán."),
        ("NS-04", "Dữ liệu quản lý thuế nội địa", "Nguồn dữ liệu cốt lõi để xác định nghĩa vụ thuế và đánh giá rủi ro thất thu."),
        ("BS-01", "Dữ liệu thu NSNN hạch toán tại Kho bạc Nhà nước", "Nguồn dữ liệu độc lập để xác minh số phải nộp, số đã nộp và hạch toán NSNN."),
        ("NS-09", "Dữ liệu hóa đơn điện tử", "Khối lượng lớn, cần kết nối trực tiếp để rà soát đồng bộ và nhận diện hóa đơn rủi ro."),
        ("NS-05", "Dữ liệu thu từ hoạt động xuất khẩu, nhập khẩu", "Địa bàn có lưu lượng xuất nhập khẩu lớn, cần dữ liệu chi tiết để đánh giá mã HS, trị giá và nghĩa vụ thuế."),
        ("CĐ-03", "Dữ liệu nghĩa vụ tài chính về đất đai", "Cần liên thông dữ liệu đất đai, thuế và Kho bạc để kiểm soát nghĩa vụ tài chính."),
        ("DN-01", "Dữ liệu đăng ký doanh nghiệp, vốn điều lệ, người đại diện", "Thông tin quan trọng phục vụ khai thác dữ liệu kiểm toán."),
        ("DN-02", "Báo cáo tài chính riêng, hợp nhất và báo cáo kiểm toán độc lập", "Thông tin quan trọng phục vụ khai thác dữ liệu kiểm toán."),
        ("DN-03", "Dữ liệu vốn nhà nước đầu tư tại doanh nghiệp; cổ tức, lợi nhuận nộp ngân sách", "Thông tin quan trọng phục vụ khai thác dữ liệu kiểm toán."),
        ("DN-04", "Thông tin công bố của doanh nghiệp niêm yết, công ty đại chúng", "Thông tin quan trọng phục vụ khai thác dữ liệu kiểm toán."),
        ("DN-05", "Dữ liệu dư nợ, phân loại nợ và dự phòng rủi ro", "Thông tin quan trọng phục vụ khai thác dữ liệu kiểm toán."),
        ("DN-06", "Dữ liệu kê khai, nộp, hoàn thuế của doanh nghiệp", "Thông tin quan trọng phục vụ khai thác dữ liệu kiểm toán."),
    ]
    priorities = [
        {"rank": rank, "code": code, "name": name, "reason": reason, "unit": "KV6"}
        for rank, (code, name, reason) in enumerate(priority_specs, 1)
    ]
    return details, priorities


def copy_files(unit, folder):
    destination = FILES_OUT / unit
    destination.mkdir(parents=True, exist_ok=True)
    result = []
    for path in sorted(folder.rglob("*")):
        if not path.is_file() or path.name.startswith("~$"):
            continue
        target = destination / path.name
        source_stat = path.stat()
        if not target.exists() or target.stat().st_size != source_stat.st_size:
            shutil.copy2(path, target)
        result.append({
            "name": path.name,
            "download": f"files/{unit}/{path.name}",
            "type": path.suffix.lower().lstrip(".").upper(),
            "size": source_stat.st_size,
            "updated": datetime.fromtimestamp(source_stat.st_mtime).isoformat(timespec="minutes"),
        })
    return result


def main():
    DIST.mkdir(parents=True, exist_ok=True)
    if PREVIEW_ONLY:
        for asset in ("index.html", "styles.css", "app.js"):
            shutil.copy2(PROJECT / "dist" / asset, DIST / asset)
    FILES_OUT.mkdir(parents=True, exist_ok=True)

    kv4_details, kv4_priorities = read_kv4(ROOT / "KV4" / "KTNN KV4 _CV998.pdf")
    kv7_details, kv7_priorities = read_excel(
        ROOT / "KV7" / "Files_30-9-2026_09-28-05" / "1. Phụ lục CV_998_CNTT.xlsx",
        "KV7", "Biểu 1", 7, "Biểu 2", 6,
    )
    kv8_details, kv8_priorities = read_excel(
        ROOT / "KV8" / "Phuluc_CV998_Layykien.xlsx",
        "KV8", "BIEU 01", 7, "BIEU 02", 7,
    )
    k12_details, k12_priorities = read_excel(
        ROOT / "K12" / "KV12. PL nguon du lieu phuc vu hoat dong kiem toan.xlsx",
        "KV12", "PL01", 8, "PL02", 6,
    )
    cn5_details, cn5_priorities = read_docx_forms(
        ROOT / "CN5" / "Phu_bieu_01_02_03_nguon_du_lieu_KTNN_CNV.docx", "CN5"
    )
    kv3_details, kv3_priorities = read_pdf_forms(
        ROOT / "KV3" / "Phu_luc xác định các nguồn dữ liệu phục vụ hoạt động kiểm toán.pdf",
        "KV3", range(1, 3), [3],
    )
    kv9_details, kv9_priorities = read_docx_forms(
        ROOT / "KV9" / "_30.9.2026_KVIX_Bao cao ve v.v xac dinh nguon du lieu theo CV 998 cua Cuc CNTT.docx",
        "KV9", 2, 3,
    )
    kv10_details, kv10_priorities = read_pdf_forms(
        ROOT / "KV10" / "771.KVX-TH. Bieu 01_ 02 Du lieu phu vu kiem toan.pdf",
        "KV10", range(1, 10), [10, 11],
    )
    cn4_details, cn4_priorities = read_pdf_forms(
        ROOT / "CN4" / "KTNN CNIV. Xac dinh nguon du lieu phuc vu kiem toan 30.9.pdf",
        "CN4", range(2, 8), [8, 9], merge_continuations=True,
    )
    kv6_details, kv6_priorities = read_kv6_scan()
    kv11_details, kv11_priorities = read_pdf_forms(
        ROOT / "KV11" / "1.1. Biểu 01_ 02_ 03 kèm theo.pdf",
        "KV11", range(1, 13), [13, 14], merge_continuations=True,
    )
    vth_details, vth_priorities = read_pdf_forms(
        ROOT / "Vụ TH" / "Phiếu TDCV gửi Cục CNTT về danh mục nguồn dữ liệu.pdf",
        "VTH", [2, 3], [4], merge_continuations=True,
    )
    vcs_details, vcs_priorities = [], []
    cnib_details, cnib_priorities = read_pdf_forms(
        ROOT / "CN Ib" / "CV gửi xác định nguồn dữ liệu-Cục CNTT.pdf",
        "CNIb", range(2, 6), [], merge_continuations=True,
    )
    cn3_details, cn3_priorities = read_pdf_forms(
        ROOT / "CN III" / "2.in 30.9. lại. CV nguồn dữ liệu.pdf",
        "CN3", range(1, 13), [13, 14], merge_continuations=True,
    )
    cn6_details, cn6_priorities = read_pdf_forms(
        ROOT / "CN VI" / "CN VI Cong van tra loi CV 998_CNTT_UDDLS.pdf",
        "CN6", range(3, 32), [32, 33], merge_continuations=True,
    )
    vp_details, vp_priorities = [], []
    thanh_tra_details = [
        to_record([
            "1", "CH-02", "07", "B2", "Kết quả thanh tra, kiểm tra",
            "Kết luận thanh tra, kiểm tra, giám sát của các cơ quan có liên quan",
            "", "Cơ quan thanh tra, kiểm tra, giám sát", "4", "1", "", "",
        ], "TTRA"),
        to_record([
            "2", "CH-02", "07", "B6", "Kết quả thanh tra, kiểm tra",
            "Kết luận thanh tra, kiểm tra, giám sát của các cơ quan có liên quan",
            "", "Cơ quan thanh tra, kiểm tra, giám sát", "4", "1", "", "",
        ], "TTRA"),
    ]
    thanh_tra_priorities = []

    details = (
        vp_details + vth_details + vcs_details + thanh_tra_details + kv3_details + kv4_details
        + kv6_details + kv7_details + kv8_details + kv9_details + kv10_details + kv11_details
        + k12_details + cnib_details + cn3_details + cn4_details + cn5_details + cn6_details
    )
    priorities = (
        vp_priorities + vth_priorities + vcs_priorities + thanh_tra_priorities + kv3_priorities
        + kv4_priorities + kv6_priorities + kv7_priorities + kv8_priorities + kv9_priorities
        + kv10_priorities + kv11_priorities + k12_priorities + cnib_priorities + cn3_priorities + cn4_priorities + cn5_priorities + cn6_priorities
    )

    file_folders = {
        "CNIb": ROOT / "CN Ib",
        "CN3": ROOT / "CN III",
        "CN6": ROOT / "CN VI",
        "VP": ROOT / "VP KTNN",
        "VTH": ROOT / "Vụ TH",
        "VCS": ROOT / "Vụ Chính sách KTNN",
        "TTRA": ROOT / "Thanh tra",
        "KV4": ROOT / "KV4",
        "KV6": ROOT / "KV6",
        "KV7": ROOT / "KV7",
        "KV8": ROOT / "KV8",
        "KV3": ROOT / "KV3",
        "KV9": ROOT / "KV9",
        "KV10": ROOT / "KV10",
        "KV11": ROOT / "KV11",
        "KV12": ROOT / "K12",
        "CN4": ROOT / "CN4",
        "CN5": ROOT / "CN5",
    }
    files = {unit: copy_files(unit, folder) for unit, folder in file_folders.items()}
    request_files = copy_files("yeu-cau-goc", ROOT / "CV yêu cầu gốc")

    unit_defs = [
        ("CNIb", "Kiểm toán nhà nước chuyên ngành Ib", cnib_details, cnib_priorities, "Có 14 dòng danh mục; chưa kèm danh sách ưu tiên và báo cáo khó khăn, vướng mắc, kiến nghị.", "Đầy đủ", "Chưa có", "Chưa có"),
        ("CN3", "Kiểm toán nhà nước chuyên ngành III", cn3_details, cn3_priorities, "Có đủ 3 biểu; danh mục ngân sách, đầu tư, doanh nghiệp và các nguồn bổ sung.", "Đầy đủ", "Đầy đủ", "Đầy đủ"),
        ("CN6", "Kiểm toán nhà nước chuyên ngành VI", cn6_details, cn6_priorities, "Có đủ 3 biểu; danh mục ngân hàng, bảo hiểm, bảo hiểm xã hội và chứng khoán.", "Đầy đủ", "Đầy đủ", "Đầy đủ"),
        ("VP", "Văn phòng Kiểm toán nhà nước", vp_details, vp_priorities, "Văn bản số 299/VP-TKTH xác nhận không phát sinh nội dung theo yêu cầu của Công văn 998.", "Không phát sinh", "Không phát sinh", "Không phát sinh"),
        ("VTH", "Vụ Tổng hợp", vth_details, vth_priorities, "Có đủ 3 biểu; gồm 10 nguồn phục vụ tổng hợp, lập kế hoạch và báo cáo toàn Ngành.", "Đầy đủ", "Đầy đủ", "Đầy đủ"),
        ("VCS", "Vụ Chính sách kiểm toán nhà nước", vcs_details, vcs_priorities, "Phiếu trao đổi chưa kèm Biểu 01, 02, 03; đề nghị tổng hợp từ các đơn vị có chức năng và đơn vị chủ trì kiểm toán.", "Chưa có", "Chưa có", "Chưa có"),
        ("TTRA", "Thanh tra Kiểm toán nhà nước", thanh_tra_details, thanh_tra_priorities, "Biểu 01 kê 2 dòng CH-02; không đề xuất nguồn ưu tiên và không nêu khó khăn, vướng mắc hoặc kiến nghị.", "Đầy đủ", "Không phát sinh", "Không phát sinh"),
        ("KV3", "Kiểm toán nhà nước khu vực III", kv3_details, kv3_priorities, "Có đủ 3 biểu; tập trung ngân sách địa phương và đầu tư công.", "Đầy đủ", "Đầy đủ", "Đầy đủ"),
        ("KV4", "Kiểm toán nhà nước khu vực IV", kv4_details, kv4_priorities, "Một PDF 16 trang; tách nội dung theo 4 lĩnh vực kiểm toán.", "Đầy đủ", "Đầy đủ", "Một phần"),
        ("KV6", "Kiểm toán nhà nước khu vực VI", kv6_details, kv6_priorities, "Có đủ 3 biểu theo 4 nhóm lĩnh vực; hồ sơ scan gồm 40 nguồn và 20 đề xuất ưu tiên.", "Đầy đủ", "Đầy đủ", "Đầy đủ"),
        ("KV7", "Kiểm toán nhà nước khu vực VII", kv7_details, kv7_priorities, "Văn bản trả lời và phụ lục Excel; tập trung ngân sách địa phương, thuế, hải quan và đầu tư.", "Đầy đủ", "Đầy đủ", "Chưa có"),
        ("KV8", "Kiểm toán nhà nước khu vực VIII", kv8_details, kv8_priorities, "Có đủ 3 biểu; phạm vi mã nguồn rộng nhất trong các hồ sơ hiện có.", "Đầy đủ", "Đầy đủ", "Đầy đủ"),
        ("KV9", "Kiểm toán nhà nước khu vực IX", kv9_details, kv9_priorities, "Có đủ 3 biểu; danh mục bao quát thu ngân sách, đất đai, doanh nghiệp và đầu tư.", "Đầy đủ", "Đầy đủ", "Đầy đủ"),
        ("KV10", "Kiểm toán nhà nước khu vực X", kv10_details, kv10_priorities, "Có đủ 3 biểu; danh mục chi tiết theo ngân sách, đầu tư và doanh nghiệp.", "Đầy đủ", "Đầy đủ", "Đầy đủ"),
        ("KV11", "Kiểm toán nhà nước khu vực XI", kv11_details, kv11_priorities, "Có đủ 3 biểu; gồm 38 dòng về ngân sách, đầu tư, doanh nghiệp và các nguồn bổ sung.", "Đầy đủ", "Đầy đủ", "Đầy đủ"),
        ("KV12", "Kiểm toán nhà nước khu vực XII", k12_details, k12_priorities, "Văn bản trả lời và phụ lục Excel; nhiều nguồn thuế, TABMIS, đầu tư và dữ liệu bổ sung.", "Đầy đủ", "Đầy đủ", "Chưa có"),
        ("CN4", "Kiểm toán nhà nước chuyên ngành IV", cn4_details, cn4_priorities, "Có đủ 3 biểu; gồm 17 nguồn dữ liệu về ngân sách, đầu tư, doanh nghiệp và kết quả kiểm tra.", "Đầy đủ", "Đầy đủ", "Đầy đủ"),
        ("CN5", "Kiểm toán nhà nước chuyên ngành V", cn5_details, cn5_priorities, "Có đủ 3 biểu; tập trung vào hai Bộ, doanh nghiệp nhà nước và các dự án đầu tư thuộc phạm vi kiểm toán.", "Đầy đủ", "Đầy đủ", "Đầy đủ"),
    ]
    roster_order = {code: index for index, (code, _, _) in enumerate(UNIT_ROSTER)}
    unit_defs.sort(key=lambda unit: roster_order[unit[0]])
    units = []
    for code, name, unit_details, unit_priorities, note, form1, form2, form3 in unit_defs:
        units.append({
            "code": code,
            "name": name,
            "detail_count": len(unit_details),
            "priority_count": len(unit_priorities),
            "form1": form1,
            "form2": form2,
            "form3": form3,
            "note": note,
            "files": files[code],
        })

    received_codes = {unit["code"] for unit in units}
    pending_units = [
        {"code": code, "name": name, "group": group}
        for code, name, group in UNIT_ROSTER if code not in received_codes
    ]

    priority_presence = Counter()
    for unit, *_ in unit_defs:
        codes = {
            code.strip()
            for p in priorities if p["unit"] == unit
            for code in re.split(r"[,;/]+", re.sub(r"([A-ZĐ])-\s+(\d)", r"\1-\2", p["code"], flags=re.IGNORECASE))
            if code.strip() and not re.fullmatch(r"BS-\d+", code.strip(), flags=re.IGNORECASE)
        }
        priority_presence.update(codes)
    bs_matches = {}
    for priority in priorities:
        code = priority["code"].upper()
        if not re.fullmatch(r"BS-\d+", code):
            continue
        key = (code, content_key(priority["name"]))
        bs_matches.setdefault(key, set()).add(priority["unit"])
    for (code, _), matched_units in bs_matches.items():
        if len(matched_units) >= 2:
            priority_presence[code] = max(priority_presence[code], len(matched_units))
    common = [
        {"code": code, "units": count}
        for code, count in sorted(priority_presence.items(), key=lambda item: (-item[1], item[0]))
    ][:12]

    issues = [
        {"severity": "Trung bình", "unit": "CNIb", "title": "Chưa có Biểu số 02 và 03", "detail": "Hồ sơ hiện có danh mục 14 dòng; hai trang cuối trống, chưa có danh sách nguồn ưu tiên và báo cáo khó khăn, vướng mắc, kiến nghị."},
        {"severity": "Trung bình", "unit": "CN3", "title": "Số thứ tự 46 bị lặp", "detail": "Hai dòng cuối Biểu 01 cùng đánh số 46 (BS-14, BS-15); tổng thực tế là 47 dòng, cần chuẩn hóa số thứ tự."},
        {"severity": "Cao", "unit": "KV8", "title": "Mã BS-01 chưa thống nhất", "detail": "Biểu 02 xếp BS-01 là báo cáo nguồn và nhu cầu cải cách tiền lương, trong khi Biểu 01 dùng BS-01 cho giấy phép khai thác khoáng sản; báo cáo cải cách tiền lương mang mã BS-02."},
        {"severity": "Cao", "unit": "KV4", "title": "Danh sách ưu tiên vượt giới hạn", "detail": "Hồ sơ nêu 19 ưu tiên theo 4 nhóm lĩnh vực, trong khi yêu cầu gốc quy định tối đa 10 nguồn cho một đơn vị."},
        {"severity": "Cao", "unit": "KV7", "title": "Danh sách ưu tiên vượt giới hạn", "detail": "Biểu 02 có 20 nguồn ưu tiên; cần chọn lại tối đa 10 nguồn và sắp xếp theo mức độ ưu tiên giảm dần."},
        {"severity": "Cao", "unit": "KV12", "title": "Danh sách ưu tiên vượt giới hạn", "detail": "Biểu 02 có 28 nguồn ưu tiên; cần chọn lại tối đa 10 nguồn theo yêu cầu của Công văn."},
        {"severity": "Cao", "unit": "KV9", "title": "Danh sách ưu tiên vượt giới hạn", "detail": "Biểu 02 có 24 nguồn ưu tiên; cần chọn lại tối đa 10 nguồn và sắp xếp theo mức độ ưu tiên giảm dần."},
        {"severity": "Cao", "unit": "KV3", "title": "Mã BS-01 chưa thống nhất", "detail": "Biểu 01 dùng BS-01 cho cả hồ sơ lập dự toán ngân sách và dữ liệu kết quả thực hiện kết luận, kiến nghị kiểm toán; cần tách mã trước khi tổng hợp toàn ngành."},
        {"severity": "Cao", "unit": "KV6", "title": "Danh sách ưu tiên vượt giới hạn", "detail": "Bốn biểu ưu tiên theo từng nhóm lĩnh vực có tổng cộng 20 nguồn; cần chọn lại tối đa 10 nguồn cho toàn đơn vị theo yêu cầu của Công văn."},
        {"severity": "Trung bình", "unit": "KV8", "title": "Mã bổ sung và số thứ tự bị lặp", "detail": "Công văn nêu 1 nguồn bổ sung nhưng phụ lục xuất hiện BS-01 đến BS-08 và nhiều đoạn đánh lại số thứ tự; cần hợp nhất trước khi tổng hợp toàn ngành."},
        {"severity": "Trung bình", "unit": "KV7", "title": "Lý do ưu tiên còn chung chung", "detail": "Nhiều dòng chỉ nêu phục vụ khảo sát/lập kế hoạch, chưa chỉ rõ hậu quả nếu thiếu dữ liệu hoặc dữ liệu giữ nguyên hình thức hiện tại."},
        {"severity": "Trung bình", "unit": "KV7", "title": "Thiếu Biểu số 03 và email người lập", "detail": "Hồ sơ hiện có 2 biểu; trường email người lập để trống và chưa có báo cáo khó khăn, vướng mắc, kiến nghị."},
        {"severity": "Trung bình", "unit": "KV12", "title": "Thiếu Biểu số 03", "detail": "Hồ sơ hiện có Biểu 01 và Biểu 02 nhưng chưa có báo cáo khó khăn, vướng mắc và kiến nghị."},
        {"severity": "Trung bình", "unit": "VCS", "title": "Chưa có Biểu số 01, 02 và 03", "detail": "Phiếu trao đổi của Vụ Chính sách kiểm toán nhà nước chưa kèm danh mục nguồn dữ liệu, danh sách ưu tiên hoặc báo cáo khó khăn, vướng mắc, kiến nghị."},
        {"severity": "Trung bình", "unit": "KV11", "title": "Số thứ tự cuối danh mục chưa liên tục", "detail": "Các dòng nguồn bổ sung cuối Biểu 01 được đánh số 35, 37, 36, 37; cần chuẩn hóa lại số thứ tự trước khi chốt danh mục toàn Ngành."},
        {"severity": "Trung bình", "unit": "KV11", "title": "Mã hình thức dữ liệu ngoài danh mục chuẩn", "detail": "Dòng DN-05 ghi hình thức dữ liệu 1, 2, 6 nhưng phụ lục hướng dẫn chỉ quy định mã 01 đến 05; giao diện đã diễn giải mã 06 là chưa xác định và cần đơn vị xác nhận lại."},
        {"severity": "Theo dõi", "unit": "KV4", "title": "Biểu số 03 chưa bao phủ đủ 4 lĩnh vực", "detail": "Có báo cáo vướng mắc cho đầu tư, chi NSNN và thu NSNN; chưa thấy phần tương ứng cho doanh nghiệp và tổ chức tài chính, ngân hàng."},
    ]

    data = {
        "meta": {
            "title": "Tổng hợp nguồn dữ liệu phục vụ hoạt động kiểm toán",
            "reference": "Công văn 998/CNTT-UDDLS ngày 24/9/2026",
            "updated": datetime.now().isoformat(timespec="minutes"),
            "unit_count": len(units),
            "detail_count": len(details),
            "priority_count": len(priorities),
            "file_count": sum(len(v) for v in files.values()),
        },
        "request": {
            "objective": "Xây dựng danh mục dữ liệu ưu tiên, kế hoạch kết nối, chia sẻ dữ liệu và kho dữ liệu dùng chung của Ngành; giảm việc yêu cầu đơn vị được kiểm toán cung cấp lại dữ liệu đã có.",
            "required": [
                "Biểu 01: Danh mục nguồn dữ liệu gắn với yêu cầu nghiệp vụ cụ thể",
                "Biểu 02: Tối đa 10 nguồn ưu tiên kết nối, khai thác trước",
                "Biểu 03: Khó khăn, vướng mắc và kiến nghị theo từng mã nguồn",
                "Yêu cầu chất lượng: đúng, đủ, sạch, sống",
                "Nêu cơ quan quản lý, phương thức, tần suất, căn cứ và điều kiện bảo mật",
            ],
            "deadline": "30/9/2026",
            "files": request_files,
        },
        "units": units,
        "pending_units": pending_units,
        "unit_roster": [{"code": code, "name": name, "group": group} for code, name, group in UNIT_ROSTER],
        "details": details,
        "priorities": priorities,
        "common_priorities": common,
        "issues": issues,
    }
    (DIST / "data.js").write_text("window.DASHBOARD_DATA = " + json.dumps(data, ensure_ascii=False) + ";\n", encoding="utf-8")
    print(json.dumps(data["meta"], ensure_ascii=True))


if __name__ == "__main__":
    main()
